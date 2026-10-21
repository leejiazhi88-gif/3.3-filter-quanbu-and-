"""Prepare name-matched note payloads from read-only Feishu XLSX exports.

Exports and payloads contain personal notes and must stay under ignored work/.
Apply via Feishu API, then re-export and compare before updating local style map.
"""
import argparse
import json
from pathlib import Path
import openpyxl


def prepare(source_file, source_title, target_file, target_title, out_dir):
    source = openpyxl.load_workbook(source_file, rich_text=True)[source_title]
    target = openpyxl.load_workbook(target_file, rich_text=True)[target_title]
    lookup = {}
    for row in range(2, source.max_row + 1):
        name = source.cell(row, 1).value
        if not name:
            continue
        assert name not in lookup, f'Duplicate source name: {name}'
        lookup[name] = source.cell(row, 2)
    values, styles, matched, unmatched = [], [], [], []
    for row in range(2, target.max_row + 1):
        name = target.cell(row, 1).value
        if not name:
            continue
        if name not in lookup:
            unmatched.append(name)
            continue
        cell = lookup[name]
        target_cell = target.cell(row, 2)
        assert target_cell.font.name == 'Calibri' and target_cell.number_format == 'General'
        assert all(not getattr(target_cell.border, edge).style for edge in ('left','right','top','bottom'))
        # Fail closed on richer styles that require a different copier.
        assert cell.value is None or isinstance(cell.value, str)
        assert not cell.hyperlink and not cell.comment
        assert cell.font.name == 'Calibri' and cell.font.sz == 10
        assert not cell.font.b and not cell.font.i and not cell.font.u and not cell.font.strike
        assert all(not getattr(cell.border, edge).style for edge in ('left','right','top','bottom'))
        assert cell.alignment.horizontal is None and cell.alignment.vertical == 'center'
        assert not cell.alignment.wrapText and not cell.alignment.textRotation
        assert cell.number_format == 'General'
        assert cell.fill.fgColor.type == 'rgb' and cell.font.color.type == 'rgb'
        style = {
            'backColor': '#' + cell.fill.fgColor.rgb[-6:],
            'foreColor': '#' + cell.font.color.rgb[-6:],
            'font': {'bold': False, 'italic': False, 'fontSize': '10pt/1.5'},
            'textDecoration': 0, 'hAlign': 0, 'vAlign': 1,
            'wrapStrategy': 0,
        }
        address = f'Wc32bt!B{row}:B{row}'
        values.append({'range': address, 'values': [[cell.value or '']]})
        styles.append({'ranges': [address], 'style': style})
        matched.append({'name': name, 'source_row': cell.row, 'target_row': row,
                        'code': target.cell(row, 5).value, 'style': style})
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    for name, data in [('values.json', values), ('styles.json', styles),
                       ('matches.json', matched), ('unmatched.json', unmatched)]:
        (out / name).write_text(json.dumps(data, ensure_ascii=False, indent=2))
    print(json.dumps({'matched': len(matched), 'unmatched': unmatched}, ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    for name in ('source_file','source_title','target_file','target_title','out_dir'):
        parser.add_argument(name)
    args = parser.parse_args()
    prepare(**vars(args))
